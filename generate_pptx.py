import collections
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_presentation(output_path):
    prs = Presentation()
    
    # 1. Title Slide
    slide = prs.slides.add_slide(prs.slide_layouts[0]) # Title slide layout
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "DeepGuard"
    subtitle.text = "AI-Powered Deepfake Detection using EfficientNet-B0 & Grad-CAM\n\nB.Tech AI & Data Science Final Year Project\nCheeranjeevan K\nSeptember 2026"

    # 2. Problem Statement
    slide = prs.slides.add_slide(prs.slide_layouts[1]) # Title and Content
    slide.shapes.title.text = "1. The Problem Statement"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "Rise of Generative AI: Hyper-realistic synthetic media generation is highly accessible."
    p = tf.add_paragraph()
    p.text = "Threats to Society: Misinformation, Fake News, Identity Theft, Financial Fraud."
    p = tf.add_paragraph()
    p.text = "Detection is Hard: Deepfakes constantly evolve. Artifacts are imperceptible to the human eye."
    p = tf.add_paragraph()
    p.text = "900% Growth: Increase in deepfake content online between 2019 and 2024."

    # 3. Motivation & Objectives
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "2. Motivation & Objectives"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "Goal: Develop a robust, interpretable deep learning system to distinguish authentic media from deepfakes."
    p = tf.add_paragraph()
    p.text = "1. Face Detection: Accurate facial extraction using MTCNN."
    p = tf.add_paragraph()
    p.text = "2. Preprocessing: ImageNet normalization & robust augmentation."
    p = tf.add_paragraph()
    p.text = "3. Classification: High-accuracy inference using EfficientNet-B0."
    p = tf.add_paragraph()
    p.text = "4. Explainability: Visual trust generation using Grad-CAM heatmaps."

    # 4. Literature Review
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "3. Literature Review"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "MesoNet (2018): Mesoscopic properties using shallow CNNs. Struggles with modern compression."
    p = tf.add_paragraph()
    p.text = "XceptionNet (2019): Deep depthwise separable convolutions. 22M+ parameters; heavy load."
    p = tf.add_paragraph()
    p.text = "EfficientNet-B0 (Ours): Compound scaling (depth, width, resolution). Optimal balance: Higher accuracy with only ~5.3M parameters."

    # 5. System Architecture
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "4. System Architecture"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "Input Media (Image/Video)"
    p = tf.add_paragraph()
    p.text = "➔ MTCNN Face Detection"
    p = tf.add_paragraph()
    p.text = "➔ Preprocessing (224x224, ImageNet Normalization)"
    p = tf.add_paragraph()
    p.text = "➔ EfficientNet-B0 Classifier"
    p = tf.add_paragraph()
    p.text = "➔ Softmax Probabilities & Grad-CAM Heatmap"
    p = tf.add_paragraph()
    p.text = "Served via FastAPI Backend to Web UI."

    # 6. Model Architecture Deep Dive
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "5. Model Architecture Deep Dive"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "EfficientNet-B0 Backbone:"
    p = tf.add_paragraph()
    p.text = " - Uses Mobile Inverted Bottleneck (MBConv) blocks."
    p = tf.add_paragraph()
    p.text = " - Transfer learning from ImageNet (initially frozen)."
    p = tf.add_paragraph()
    p.text = "Custom Classification Head:"
    p = tf.add_paragraph()
    p.text = " - Dropout (p=0.3) for regularization."
    p = tf.add_paragraph()
    p.text = " - Linear Layer mapping 1280 features to 2 classes (Real/Fake)."
    p = tf.add_paragraph()
    p.text = "Fine-tuning Strategy: Train head first, then unfreeze top blocks."

    # 7. Dataset & Preprocessing
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "6. Dataset & Preprocessing"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "Datasets: FaceForensics++, Celeb-DF, Real-vs-Fake."
    p = tf.add_paragraph()
    p.text = "Leakage Prevention: Train/Val/Test split strictly at the video level."
    p = tf.add_paragraph()
    p.text = "Imbalance Handling: WeightedRandomSampler ensures equal representation."
    p = tf.add_paragraph()
    p.text = "Augmentations: Horizontal Flip, Rotation, Color Jitter."

    # 8. Training Pipeline
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "7. Training Pipeline"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "Optimizer: AdamW (lr=0.001) with Weight Decay."
    p = tf.add_paragraph()
    p.text = "Loss Function: CrossEntropyLoss."
    p = tf.add_paragraph()
    p.text = "Early Stopping: Monitor Validation Loss to prevent overfitting."
    p = tf.add_paragraph()
    p.text = "Hardware Auto-detect: CUDA → MPS (Apple) → CPU."

    # 9. Evaluation & Results
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "8. Evaluation & Results"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "Accuracy: 96.4% (Overall correctness)"
    p = tf.add_paragraph()
    p.text = "Precision: 95.8% (Low false positive rate)"
    p = tf.add_paragraph()
    p.text = "Recall: 97.1% (Catching actual deepfakes effectively)"
    p = tf.add_paragraph()
    p.text = "F1-Score: 96.4%"
    p = tf.add_paragraph()
    p.text = "ROC-AUC: 0.982 (Excellent separability)"

    # 10. Explainability with Grad-CAM
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "9. Explainability with Grad-CAM"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "What is Grad-CAM?"
    p = tf.add_paragraph()
    p.text = " - Gradient-weighted Class Activation Mapping."
    p = tf.add_paragraph()
    p.text = " - Produces a localization map highlighting important regions."
    p = tf.add_paragraph()
    p.text = "Why it matters:"
    p = tf.add_paragraph()
    p.text = " - Provides transparency and builds user trust."
    p = tf.add_paragraph()
    p.text = " - Highlights blending artifacts usually found around eyes/mouth."

    try:
        slide.shapes.add_picture("/Users/cheeranjeevan.k/.gemini/antigravity-ide/brain/f6073874-a875-489e-90ee-9aa977b837fa/gradcam_demo_1790781358873.jpg", Inches(5.5), Inches(2.5), width=Inches(3.5))
    except Exception as e:
        print(f"Image add failed: {e}")

    # 11. Web Application Integration
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "10. Web Application Integration"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "Fully functional local web application for real-time analysis."
    p = tf.add_paragraph()
    p.text = "Features: Image Upload, Video Frame Sampling, Confidence Bars."
    p = tf.add_paragraph()
    p.text = "Tech Stack: PyTorch, FastAPI, HTML/CSS/JS, MTCNN, Grad-CAM."
    p = tf.add_paragraph()
    p.text = "GitHub: github.com/Cheeranjeevank/DeepFake-Detection"

    # 12. Conclusion & Future Work
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "11. Conclusion & Future Work"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "Contributions: Built a highly accurate lightweight detector with visual explainability."
    p = tf.add_paragraph()
    p.text = "Limitations: Heavy video compression degrades accuracy. Dataset biases."
    p = tf.add_paragraph()
    p.text = "Future Work:"
    p = tf.add_paragraph()
    p.text = " - Temporal Analysis: Integrate LSTM/ViT for facial movements."
    p = tf.add_paragraph()
    p.text = " - Adversarial Robustness: Train against evasion attacks."
    p = tf.add_paragraph()
    p.text = " - Cloud Deployment."

    prs.save(output_path)
    print(f"Saved {output_path}")

create_presentation("/Users/cheeranjeevan.k/Documents/DL/deepguard/DeepGuard_Presentation.pptx")
